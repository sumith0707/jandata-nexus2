"""
Extracts everything of interest from a .docx file, regardless of layout:
- plain paragraph text
- native Word tables (real table objects, no OCR needed)
- embedded images (may contain tables-as-pictures, need OCR)

This file has no knowledge of any specific report's structure — it works
the same way on any .docx.
"""

import os
import zipfile
from docx import Document


def extract_text_and_native_tables(docx_path: str):
    """
    Returns:
        paragraphs: list[str] — plain text content, in order
        native_tables: list[list[list[str]]] — each table as rows of cell text
    """
    doc = Document(docx_path)

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    native_tables = []
    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        native_tables.append(rows)

    return paragraphs, native_tables


def extract_embedded_images(docx_path: str, output_dir: str):
    """
    Docx files are zip archives; embedded images live under word/media/.
    Returns list of saved image file paths.
    """
    os.makedirs(output_dir, exist_ok=True)
    saved_paths = []

    with zipfile.ZipFile(docx_path) as z:
        for name in z.namelist():
            if name.startswith("word/media/"):
                filename = os.path.basename(name)
                target_path = os.path.join(output_dir, filename)
                with z.open(name) as src, open(target_path, "wb") as dst:
                    dst.write(src.read())
                saved_paths.append(target_path)

    return saved_paths


if __name__ == "__main__":
    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else "data/raw/Kharif2025SeasonalConditionason25072025Eng.docx"
    paragraphs, native_tables = extract_text_and_native_tables(path)
    images = extract_embedded_images(path, "data/processed/extracted_images")

    print(f"Paragraphs: {len(paragraphs)}")
    print(f"Native Word tables: {len(native_tables)}")
    print(f"Embedded images: {len(images)} -> {images}")
